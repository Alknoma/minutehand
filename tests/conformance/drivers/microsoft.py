"""Microsoft, driven as a Teams bot and a Microsoft 365 app drive it: the identity platform's token endpoint
(https://learn.microsoft.com/en-us/entra/identity-platform/v2-oauth2-client-creds-grant-flow), the Bot Framework
connector a bot sends with (https://learn.microsoft.com/en-us/azure/bot-service/rest-api/bot-framework-rest-connector-api-reference),
and Microsoft Graph `v1.0` for reads, files and everything a signed-in user does (https://learn.microsoft.com/en-us/graph/api/overview).

A world is one tenant, its own by `tag`: the scenario's name is the tag, and the tenant id, its domain, its
SharePoint host, the bot's app id and secret and the team are derived from that name, as the provider's README says
a service is configured with them before the run (`state.directory_of`). The agent is the bot's app: client
credentials for Graph (application permissions) and for the Bot Framework. A person signs in to the same app by the
authorization code flow with their sign-in name as `login_hint` (the provider's authorize has no browser), and acts
through Graph with the delegated token, as a person's own client does.

No Microsoft client library is a dev dependency (neither `msgraph-sdk` nor `msal`), so every call is plain httpx
through `Api.http`, as the REST references above write them, and every answer is read as they document it.
"""

from __future__ import annotations

import html
import io
import json
import re
import uuid
from collections.abc import Callable, Mapping
from datetime import UTC, datetime
from typing import Any, ClassVar
from urllib.parse import parse_qs, quote, urlsplit

import docx
import httpx

from minutehand.adapters.control.wire import Claims, CreateWorld, WorldView
from minutehand.adapters.providers.microsoft.state import Directory, directory_of
from minutehand.domain.scenario import AccessRole, DocumentKind, Seed
from tests.conformance.contract import (
    STAMP_CHANNEL,
    Api,
    ChannelSeen,
    CommentSeen,
    Documents,
    DocumentSeen,
    Driver,
    FaultCase,
    IdKind,
    MessageSeen,
    Messaging,
    PersonSeen,
    Session,
    VendorRefused,
    ok,
)

PROVIDER = "microsoft"
LOGIN = "https://login.microsoftonline.com"
GRAPH = "https://graph.microsoft.com/v1.0"
CONNECTOR = "https://smba.trafficmanager.net/teams/"
"""The `serviceUrl` Teams names on every activity it pushes a bot (the provider's README)."""
GRAPH_DEFAULT = "https://graph.microsoft.com/.default"
BOT_DEFAULT = "https://api.botframework.com/.default"
REDIRECT = "https://localhost/conformance/signed-in"
DELEGATED = (
    "openid profile offline_access User.Read User.ReadBasic.All Files.ReadWrite.All Sites.ReadWrite.All "
    "Team.ReadBasic.All Channel.ReadBasic.All ChannelMessage.Read.All ChannelMessage.Send Chat.ReadWrite"
)
"""The delegated permissions a person's client asks for (https://learn.microsoft.com/en-us/graph/permissions-reference)."""
USER_FIELDS = "id,displayName,mail,userPrincipalName,jobTitle,accountEnabled,userType"
"""`accountEnabled` and `userType` are not in a user's default properties and must be selected
(https://learn.microsoft.com/en-us/graph/api/resources/user#properties)."""
ITEM_FIELDS = "id,name,size,createdDateTime,lastModifiedDateTime,createdBy,lastModifiedBy,parentReference,file,folder"
"""What a listing selects of a driveItem: everything but `@microsoft.graph.downloadUrl`, a short-lived URL."""
USERS_PAGE = 999
MESSAGES_PAGE = 50
ITEMS_PAGE = 999
DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
OFFICE_KINDS = {
    DOCX: DocumentKind.DOCUMENT,
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": DocumentKind.SPREADSHEET,
    "application/vnd.openxmlformats-officedocument.presentationml.presentation": DocumentKind.PRESENTATION,
}
OFFICE_ENDINGS = (".docx", ".xlsx", ".pptx")
"""An Office file shows in SharePoint and Teams under its name without this ending; the ending stays on the item."""
CARD = "application/vnd.microsoft.card.adaptive"
ROLES = {AccessRole.READER: "read", AccessRole.WRITER: "write", AccessRole.ORGANIZER: "owner"}
"""driveItem invite's roles (https://learn.microsoft.com/en-us/graph/api/driveitem-invite): read, write, owner."""
GRANTED = {"read": AccessRole.READER, "write": AccessRole.WRITER, "owner": AccessRole.ORGANIZER}
REACTIONS = {
    "thumbsup": "👍",
    "+1": "👍",
    "heart": "❤️",
    "laughing": "😆",
    "joy": "😂",
    "open_mouth": "😮",
    "cry": "😢",
    "angry": "😠",
    "eyes": "👀",
    "tada": "🎉",
    "white_check_mark": "✅",
}
"""A reaction's neutral name as the Unicode character Graph's `setReaction` takes as `reactionType`
(https://learn.microsoft.com/en-us/graph/api/chatmessage-setreaction)."""
LEGACY_REACTIONS = {"like": "thumbsup", "heart": "heart", "laugh": "laughing", "surprised": "open_mouth",
                    "sad": "cry", "angry": "angry"}  # fmt: skip
"""Teams' legacy reaction types, which `chatMessageReaction.reactionType` may still answer."""
GRAPH_TIME = re.compile(r"^(\d{4}-\d{2}-\d{2}T\d{2}:\d{2}:\d{2})(?:\.(\d{1,7}))?Z$")
"""Graph's `DateTimeOffset`: ISO 8601, always UTC, `Z` (https://learn.microsoft.com/en-us/graph/api/resources/chatmessage)."""
UNSEEDED = (
    "eyJ0eXAiOiJKV1QiLCJhbGciOiJSUzI1NiJ9.eyJhdWQiOiJodHRwczovL2dyYXBoLm1pY3Jvc29mdC5jb20iLCJ0aWQiOiJ1bnNlZWRlZCJ9."
    "dW5zZWVkZWQ"
)
"""A Graph access token of the identity platform's shape (a signed JWT) that no world's sign-in minted."""

Doc = Mapping[str, Any]
"""A JSON object as Microsoft answered it, read only inside this driver."""


def directory(name: str) -> Directory:
    """The tenant a world of this name is: every id a service is configured with (the provider's README)."""
    seed = Seed.model_validate(
        {"name": name, "people": [{"key": "x", "name": "X", "email": "x@example.com", "reply": {"kind": "silent"}}]}
    )
    return directory_of(seed.starting(datetime(2026, 1, 1, tzinfo=UTC)))


def graph_time(text: str) -> datetime:
    """A Graph timestamp as aware UTC; anything not in Graph's format is refused."""
    found = GRAPH_TIME.match(text)
    if found is None:
        raise ValueError(f"{text!r} is not a Graph DateTimeOffset")
    fraction = (found.group(2) or "").ljust(6, "0")[:6]
    return datetime.strptime(f"{found.group(1)}.{fraction}", "%Y-%m-%dT%H:%M:%S.%f").replace(tzinfo=UTC)


def _json(what: str, response: httpx.Response) -> Doc:
    ok(what, response)
    body = response.json()
    if not isinstance(body, dict):
        raise VendorRefused(f"{what}: the answer is not a JSON object", response)
    return body


def _given(found: Doc, key: str) -> bool:
    """The answer carries `key` with a value that is not null, empty or false."""
    return key in found and bool(found[key])


def _deleted(item: Doc) -> bool:
    """A driveItem carrying the `deleted` facet, which may be an empty object (driveitem-delta)."""
    return "deleted" in item and item["deleted"] is not None


def _expect(what: str, response: httpx.Response, status: int) -> httpx.Response:
    """The response, refused unless it carries the status the vendor documents for the operation."""
    if response.status_code != status:
        raise VendorRefused(f"{what}: documented {status}, answered", response)
    return response


def _text(body: Doc) -> str:
    """A chatMessage's `body` as text: an `html` body has its markup removed and its entities read."""
    content = str(body["content"]) if _given(body, "content") else ""
    if "contentType" in body and body["contentType"] == "html":
        return html.unescape(re.sub(r"<[^>]+>", "", content))
    return content


def _docx(text: str) -> bytes:
    """A Word document holding `text`, one paragraph per line."""
    made = docx.Document()
    for line in text.split("\n"):
        made.add_paragraph(line)
    out = io.BytesIO()
    made.save(out)
    return out.getvalue()


def _docx_text(content: bytes) -> str:
    return "\n".join(p.text for p in docx.Document(io.BytesIO(content)).paragraphs)


def _segments(path: str) -> list[str]:
    return [p for p in path.split("/") if p]


def _quoted(path: list[str]) -> str:
    return "/".join(quote(p, safe="") for p in path)


def _title(name: str) -> str:
    lower = name.lower()
    return next((name[: -len(e)] for e in OFFICE_ENDINGS if lower.endswith(e)), name)


class SignIn:
    """The identity platform's token endpoint for one tenant, through the recording client."""

    def __init__(self, http: httpx.Client, tenant: Directory) -> None:
        self._http = http
        self._tenant = tenant
        self.minted: list[str] = []

    def _token(self, form: Mapping[str, str]) -> Doc:
        answered = _json(
            f"sign in ({form['grant_type']})",
            self._http.post(f"{LOGIN}/{self._tenant.tenant_id}/oauth2/v2.0/token", data=dict(form)),
        )
        for field in ("access_token", "refresh_token", "id_token"):
            if field in answered:
                self.minted.append(str(answered[field]))
        return answered

    def app(self, scope: str) -> str:
        """Client credentials for the bot's app (v2-oauth2-client-creds-grant-flow)."""
        return str(
            self._token(
                {
                    "grant_type": "client_credentials",
                    "client_id": self._tenant.bot_app_id,
                    "client_secret": self._tenant.bot_app_secret,
                    "scope": scope,
                }
            )["access_token"]
        )

    def user(self, login: str) -> str:
        """The authorization code flow (v2-oauth2-auth-code-flow) for the user signing in as `login`."""
        authorized = self._http.get(
            f"{LOGIN}/{self._tenant.tenant_id}/oauth2/v2.0/authorize",
            params={"client_id": self._tenant.bot_app_id, "response_type": "code", "redirect_uri": REDIRECT,
                    "response_mode": "query", "scope": DELEGATED, "login_hint": login, "state": "conformance"},
        )  # fmt: skip
        _expect("authorize", authorized, 302)
        query = parse_qs(urlsplit(authorized.headers["location"]).query)
        if "code" not in query:
            raise VendorRefused("authorize: the redirect carries no code", authorized)
        code = query["code"][0]
        self.minted.append(code)
        return str(
            self._token(
                {
                    "grant_type": "authorization_code",
                    "client_id": self._tenant.bot_app_id,
                    "client_secret": self._tenant.bot_app_secret,
                    "code": code,
                    "redirect_uri": REDIRECT,
                    "scope": DELEGATED,
                }
            )["access_token"]
        )


class MicrosoftSession(Messaging, Documents):
    """The bot's app (`login` None): Graph with application permissions, the connector to send. A person: Graph
    with their delegated token for everything, as Teams and OneDrive clients call it."""

    def __init__(self, api: Api, tenant: Directory, login: str | None) -> None:
        self._tenant = tenant
        self._http = api.http()
        self._as_app = login is None
        signing = SignIn(self._http, tenant)
        self._graph_token = signing.app(GRAPH_DEFAULT) if login is None else signing.user(login)
        self._bot_token = signing.app(BOT_DEFAULT) if login is None else None
        self.secrets = (tenant.bot_app_secret, *signing.minted, UNSEEDED)
        self._drive: str | None = None
        self._site_name: str | None = None
        self._emails: dict[str, str | None] | None = None

    def close(self) -> None:
        self._http.close()

    # ------------------------------------------------------------------ the surfaces

    def _graph(self, method: str, path: str, **kw: Any) -> httpx.Response:
        url = path if path.startswith("https://") else f"{GRAPH}{path}"
        headers = {"Authorization": f"Bearer {self._graph_token}", **(kw.pop("headers") if "headers" in kw else {})}
        return self._http.request(method, url, headers=headers, **kw)

    def _read(self, path: str, **params: str | int) -> Doc:
        return _json(f"GET {path}", self._graph("GET", path, params=params))

    def _pages(self, path: str, **params: str | int) -> list[list[Doc]]:
        """Every page of a Graph collection, following `@odata.nextLink` (https://learn.microsoft.com/en-us/graph/paging)."""
        pages: list[list[Doc]] = []
        answered = self._read(path, **params)
        while True:
            pages.append(list(answered["value"]))
            if "@odata.nextLink" not in answered or not answered["@odata.nextLink"]:
                return pages
            following = str(answered["@odata.nextLink"])
            answered = _json(f"GET {following}", self._graph("GET", following))

    def _bot(self, method: str, path: str, body: Mapping[str, object] | None = None) -> httpx.Response:
        if self._bot_token is None:
            raise AssertionError("a person's session holds no Bot Framework token")
        return self._http.request(
            method,
            f"{CONNECTOR}v3/{path}",
            headers={"Authorization": f"Bearer {self._bot_token}"},
            json=dict(body) if body is not None else None,
        )

    # ------------------------------------------------------------------ accounts

    def whoami(self) -> str:
        if self._as_app:
            # An app is a service principal, read by its app id (https://learn.microsoft.com/en-us/graph/api/serviceprincipal-get).
            found = self._read(f"/servicePrincipals(appId='{self._tenant.bot_app_id}')", **{"$select": "displayName"})
            return str(found["displayName"])
        return str(self._read("/me", **{"$select": "userPrincipalName"})["userPrincipalName"])

    def _users(self, page_size: int = USERS_PAGE) -> list[list[Doc]]:
        return self._pages("/users", **{"$select": USER_FIELDS, "$top": page_size})

    def people(self) -> list[PersonSeen]:
        return [
            PersonSeen(
                id=str(u["id"]),
                name=str(u["displayName"]),
                email=str(u["mail"]) if _given(u, "mail") else None,
                active=not ("accountEnabled" in u and u["accountEnabled"] is False),
                guest="userType" in u and u["userType"] == "Guest",
                title=str(u["jobTitle"]) if _given(u, "jobTitle") else None,
                login=str(u["userPrincipalName"]),
            )
            for page in self._users()
            for u in page
        ]

    def _email_of(self) -> dict[str, str | None]:
        if self._emails is None:
            self._emails = {p.id: p.email for p in self.people()}
        return self._emails

    def _identity_email(self, identity: object) -> str | None:
        """The email of the user an identitySet names; None for an application's."""
        if not isinstance(identity, Mapping) or "user" not in identity or not identity["user"]:
            return None
        user = identity["user"]
        if _given(user, "email"):
            return str(user["email"])
        emails = self._email_of()
        return emails[str(user["id"])] if str(user["id"]) in emails else None

    def people_pages(self, page_size: int) -> list[list[str]]:
        return [[str(u["id"]) for u in page] for page in self._pages("/users", **{"$select": "id", "$top": page_size})]

    def unknown_credential(self) -> httpx.Response:
        """A Graph read with an access token no sign-in minted. (Client credentials for an unknown app at the
        tenant's token endpoint would be routed to the tenant's world by the tenant in the path.)"""
        return self._http.get(f"{GRAPH}/users", headers={"Authorization": f"Bearer {UNSEEDED}"})

    def observe(self) -> str:
        launch = self.channel(STAMP_CHANNEL)
        drive = self._drive_id()
        answered: list[str] = []
        for path, params in (
            ("/users", {"$select": USER_FIELDS, "$top": USERS_PAGE}),
            (f"/teams/{self._tenant.team_id}/channels", {}),
            (f"/teams/{self._tenant.team_id}/channels/{quote(launch)}/messages",
             {"$expand": "replies", "$top": MESSAGES_PAGE}),
            (f"/drives/{drive}/root/children", {"$select": ITEM_FIELDS, "$top": ITEMS_PAGE}),
        ):  # fmt: skip
            response = self._graph("GET", path, params=params)
            ok(f"GET {path}", response)
            answered.append(response.text)
        return "\n".join(answered)

    def change(self, label: str) -> None:
        self.post(self.channel(STAMP_CHANNEL), label)

    # ------------------------------------------------------------------ conversations

    def _team(self) -> str:
        return f"/teams/{self._tenant.team_id}"

    @staticmethod
    def _chat(conversation: str) -> bool:
        """A chat rather than a team's channel, by Teams' own id formats: the connector's `a:` id of a 1:1 with a
        bot, Graph's `…@unq.gbl.spaces` of a 1:1 and `…@thread.v2` of a group chat; a channel is `…@thread.tacv2`
        or `…@thread.skype` (https://learn.microsoft.com/en-us/graph/api/resources/chat)."""
        return conversation.startswith("a:") or conversation.endswith(("@unq.gbl.spaces", "@thread.v2"))

    def _graph_chat(self, conversation: str) -> str:
        """Graph's id of a chat. A bot's 1:1 named by the connector (`a:…`) is `19:{user}_{app}@unq.gbl.spaces`,
        its user read from the conversation's roster (https://learn.microsoft.com/en-us/microsoftteams/platform/bots/how-to/get-teams-context)."""
        if not conversation.startswith("a:"):
            return conversation
        roster = self._bot("GET", f"conversations/{quote(conversation, safe='')}/members")
        members = ok("the 1:1's roster", roster).json()
        if not isinstance(members, list) or len(members) != 1:
            raise VendorRefused("the 1:1's roster does not name one member", roster)
        return f"19:{members[0]['aadObjectId']}_{self._tenant.bot_app_id}@unq.gbl.spaces"

    def _messages_path(self, conversation: str) -> str:
        if self._chat(conversation):
            return f"/chats/{quote(self._graph_chat(conversation), safe='')}/messages"
        return f"{self._team()}/channels/{quote(conversation, safe='')}/messages"

    def _channels(self) -> list[Doc]:
        return [c for page in self._pages(f"{self._team()}/channels") for c in page]

    def channel(self, name: str) -> str:
        found = [
            str(c["id"])
            for page in self._pages(f"{self._team()}/channels", **{"$filter": f"displayName eq '{name}'"})
            for c in page
        ]
        if len(found) != 1:
            raise LookupError(f"Graph lists {len(found)} channels of the team named {name!r}")
        return found[0]

    def direct(self, person: str) -> str:
        if self._as_app:
            # A proactive 1:1 (https://learn.microsoft.com/en-us/microsoftteams/platform/bots/how-to/conversations/send-proactive-messages).
            made = self._bot(
                "POST",
                "conversations",
                {"bot": {"id": f"28:{self._tenant.bot_app_id}"}, "members": [{"id": person}], "isGroup": False,
                 "tenantId": self._tenant.tenant_id, "channelData": {"tenant": {"id": self._tenant.tenant_id}}},
            )  # fmt: skip
            return str(_json("create the 1:1", made)["id"])
        me = str(self._read("/me", **{"$select": "id"})["id"])

        def member(user: str) -> dict[str, object]:
            return {"@odata.type": "#microsoft.graph.aadUserConversationMember", "roles": ["owner"],
                    "user@odata.bind": f"{GRAPH}/users('{user}')"}  # fmt: skip

        made = self._graph("POST", "/chats", json={"chatType": "oneOnOne", "members": [member(me), member(person)]})
        # Create chat answers 201 with the chat, or the existing 1:1 (https://learn.microsoft.com/en-us/graph/api/chat-post).
        return str(_json("create the 1:1", _expect("create the 1:1", made, 201))["id"])

    def channels(self) -> list[ChannelSeen]:
        found: list[ChannelSeen] = []
        for c in self._channels():
            members = [
                m
                for page in self._pages(f"{self._team()}/channels/{quote(str(c['id']), safe='')}/members")
                for m in page
            ]
            found.append(
                ChannelSeen(
                    id=str(c["id"]),
                    name=str(c["displayName"]),
                    private="membershipType" in c and c["membershipType"] == "private",
                    archived="isArchived" in c and c["isArchived"] is True,
                    purpose=str(c["description"]) if _given(c, "description") else "",
                    member_emails=frozenset(str(m["email"]) for m in members if _given(m, "email")),
                )
            )
        return found

    def _send(self, conversation: str, activity: Mapping[str, object], *, reply_to: str | None = None) -> str:
        where = quote(conversation, safe="")
        path = f"conversations/{where}/activities" + (f"/{reply_to}" if reply_to is not None else "")
        return str(_json("send an activity", _expect("send", self._bot("POST", path, activity), 201))["id"])

    def _graph_post(self, conversation: str, body: Mapping[str, object], *, reply_to: str | None = None) -> str:
        path = self._messages_path(conversation)
        if reply_to is not None and not self._chat(conversation):
            path += f"/{reply_to}/replies"
        made = self._graph("POST", path, json=dict(body))
        # Send chatMessage answers 201 with the message (https://learn.microsoft.com/en-us/graph/api/chatmessage-post).
        return str(_json(f"POST {path}", _expect(f"POST {path}", made, 201))["id"])

    def post(self, channel: str, text: str) -> str:
        if self._as_app:
            return self._send(channel, {"type": "message", "text": text})
        return self._graph_post(channel, {"body": {"contentType": "text", "content": text}})

    def post_button(self, channel: str, text: str, action_id: str, label: str) -> str:
        """An Adaptive Card with one `Action.Execute` (https://learn.microsoft.com/en-us/microsoftteams/platform/task-modules-and-cards/cards/universal-actions-for-adaptive-cards/overview)."""
        card = {
            "type": "AdaptiveCard",
            "$schema": "http://adaptivecards.io/schemas/adaptive-card.json",
            "version": "1.4",
            "body": [{"type": "TextBlock", "text": text, "wrap": True}],
            "actions": [{"type": "Action.Execute", "title": label, "verb": action_id, "id": action_id}],
        }
        if self._as_app:
            return self._send(channel, {"type": "message", "attachments": [{"contentType": CARD, "content": card}]})
        attachment = str(uuid.uuid5(uuid.NAMESPACE_URL, f"{channel}/{action_id}/{text}"))
        return self._graph_post(
            channel,
            {
                "body": {"contentType": "html", "content": f'<attachment id="{attachment}"></attachment>'},
                "attachments": [{"id": attachment, "contentType": CARD, "content": json.dumps(card)}],
            },
        )

    def reply(self, channel: str, thread: str, text: str) -> str:
        if self._as_app:
            return self._send(channel, {"type": "message", "text": text, "replyToId": thread}, reply_to=thread)
        return self._graph_post(channel, {"body": {"contentType": "text", "content": text}}, reply_to=thread)

    def edit(self, channel: str, message: str, text: str) -> None:
        if self._as_app:
            where = f"conversations/{quote(channel, safe='')}/activities/{message}"
            _json("update an activity", self._bot("PUT", where, {"type": "message", "id": message, "text": text}))
            return
        path = f"{self._messages_path(channel)}/{message}"
        # Update chatMessage answers 204 (https://learn.microsoft.com/en-us/graph/api/chatmessage-update).
        _expect(
            f"PATCH {path}", self._graph("PATCH", path, json={"body": {"contentType": "text", "content": text}}), 204
        )

    def react(self, channel: str, message: str, reaction: str) -> None:
        path = f"{self._messages_path(channel)}/{message}/setReaction"
        typed = REACTIONS[reaction] if reaction in REACTIONS else reaction
        # setReaction answers 204 (https://learn.microsoft.com/en-us/graph/api/chatmessage-setreaction).
        _expect(f"POST {path}", self._graph("POST", path, json={"reactionType": typed}), 204)

    def delete_message(self, channel: str, message: str) -> None:
        if self._as_app:
            ok(
                "delete an activity",
                self._bot("DELETE", f"conversations/{quote(channel, safe='')}/activities/{message}"),
            )
            return
        path = f"{self._messages_path(channel)}/{message}/softDelete"
        # softDelete answers 204 (https://learn.microsoft.com/en-us/graph/api/chatmessage-softdelete).
        _expect(f"POST {path}", self._graph("POST", path), 204)

    def _message(self, m: Doc) -> MessageSeen:
        attachments = list(m["attachments"]) if _given(m, "attachments") else []
        reactions = list(m["reactions"]) if _given(m, "reactions") else []
        named = {v: k for k, v in REACTIONS.items() if k != "+1"}
        return MessageSeen(
            id=str(m["id"]),
            text=_text(m["body"]),
            author_email=self._identity_email(m["from"]) if "from" in m else None,
            thread_of=str(m["replyToId"]) if _given(m, "replyToId") else None,
            at=graph_time(str(m["createdDateTime"])),
            reactions=tuple(
                named[t] if t in named else LEGACY_REACTIONS[t] if t in LEGACY_REACTIONS else t
                for r in reactions
                for t in [str(r["reactionType"])]
            ),
            files=tuple(
                str(a["name"])
                for a in attachments
                if _given(a, "name") and not str(a["contentType"]).startswith("application/vnd.microsoft.card")
            ),
        )

    def history(self, channel: str) -> list[MessageSeen]:
        """A channel's roots with their replies expanded (list-channel-messages, `$expand=replies`), or a chat's
        messages; Graph answers newest first, so they are put in creation order."""
        path = self._messages_path(channel)
        expand = {} if self._chat(channel) else {"$expand": "replies"}
        every: list[Doc] = []
        for page in self._pages(path, **expand, **{"$top": MESSAGES_PAGE}):
            for root in page:
                every.append(root)
                every += list(root["replies"]) if _given(root, "replies") else []
        kept = [
            m
            for m in every
            if not _given(m, "deletedDateTime") and not ("messageType" in m and m["messageType"] != "message")
        ]
        kept.sort(key=lambda m: (graph_time(str(m["createdDateTime"])), str(m["id"])))
        return [self._message(m) for m in kept]

    def history_pages(self, channel: str, page_size: int) -> list[list[str]]:
        return [
            [str(m["id"]) for m in page] for page in self._pages(self._messages_path(channel), **{"$top": page_size})
        ]

    # ------------------------------------------------------------------ documents

    def _drive_id(self) -> str:
        """The team's document library: its site found by the team's name (site-search), then its default drive
        (https://learn.microsoft.com/en-us/graph/api/drive-get)."""
        if self._drive is None:
            team = str(self._read(self._team(), **{"$select": "displayName"})["displayName"])
            sites = [s for s in self._read("/sites", search=team)["value"] if s["displayName"] == team]
            if len(sites) != 1:
                raise LookupError(f"Graph's site search finds {len(sites)} sites named {team!r}")
            self._site_name = team
            self._drive = str(self._read(f"/sites/{sites[0]['id']}/drive", **{"$select": "id"})["id"])
        return self._drive

    def _items(self, page_size: int = ITEMS_PAGE) -> list[list[Doc]]:
        """The whole drive, from the root's `delta` with no token (https://learn.microsoft.com/en-us/graph/api/driveitem-delta)."""
        return self._pages(f"/drives/{self._drive_id()}/root/delta", **{"$top": page_size})

    def documents(self) -> list[DocumentSeen]:
        every = [i for page in self._items() for i in page if not _deleted(i)]
        folders = {str(i["id"]): i for i in every if _given(i, "folder")}

        def path_of(item: Doc) -> str | None:
            names: list[str] = []
            parent = item["parentReference"]["id"] if "id" in item["parentReference"] else None
            while parent is not None and parent in folders and "root" not in folders[parent]:
                names.append(str(folders[parent]["name"]))
                above = folders[parent]["parentReference"]
                parent = above["id"] if "id" in above else None
            return "/".join(reversed(names)) or None

        found: list[DocumentSeen] = []
        for item in every:
            if "file" not in item or item["file"] is None:
                continue
            mime = str(item["file"]["mimeType"]) if "mimeType" in item["file"] else None
            found.append(
                DocumentSeen(
                    id=str(item["id"]),
                    title=_title(str(item["name"])),
                    kind=OFFICE_KINDS[mime] if mime in OFFICE_KINDS else DocumentKind.FILE,
                    folder=path_of(item),
                    owner_email=self._identity_email(item["createdBy"]) if "createdBy" in item else None,
                    space=self._site_name,
                    shared=self._shared(str(item["id"])),
                    last_editor_email=self._identity_email(item["lastModifiedBy"])
                    if "lastModifiedBy" in item
                    else None,
                    modified=graph_time(str(item["lastModifiedDateTime"])),
                    created=graph_time(str(item["createdDateTime"])),
                    mime_type=mime,
                )
            )
        return found

    def _shared(self, item: str) -> dict[str, AccessRole]:
        """Who the item's permissions name, and the most each gives (https://learn.microsoft.com/en-us/graph/api/driveitem-list-permissions)."""
        given: dict[str, AccessRole] = {}
        order = list(AccessRole)
        for p in self._read(f"/drives/{self._drive_id()}/items/{item}/permissions")["value"]:
            to = self._identity_email(p["grantedToV2"]) if "grantedToV2" in p else None
            if to is None:
                continue
            roles = [GRANTED[r] for r in p["roles"] if r in GRANTED]
            for role in roles:
                if to not in given or order.index(role) > order.index(given[to]):
                    given[to] = role
        return given

    def _item(self, document: str) -> Doc:
        return self._read(f"/drives/{self._drive_id()}/items/{document}")

    def _folder(self, path: str | None) -> str:
        """The folder at that path, each one missing made under the last (driveitem-post-children)."""
        drive = self._drive_id()
        here = str(self._read(f"/drives/{drive}/root", **{"$select": "id"})["id"])
        walked: list[str] = []
        for name in _segments(path or ""):
            walked.append(name)
            found = self._graph("GET", f"/drives/{drive}/root:/{_quoted(walked)}", params={"$select": "id"})
            if found.status_code == 404:
                made = self._graph(
                    "POST",
                    f"/drives/{drive}/items/{here}/children",
                    json={"name": name, "folder": {}, "@microsoft.graph.conflictBehavior": "fail"},
                )
                here = str(_json(f"make the folder {name}", _expect(f"make the folder {name}", made, 201))["id"])
            else:
                here = str(_json(f"the folder {name}", found)["id"])
        return here

    def _put(self, folder: str, name: str, content: bytes, mime: str) -> str:
        """Upload by path under a folder (driveitem-put-content): 201 Created."""
        path = f"/drives/{self._drive_id()}/items/{folder}:/{quote(name, safe='')}:/content"
        made = self._graph(
            "PUT",
            path,
            content=content,
            params={"@microsoft.graph.conflictBehavior": "fail"},
            headers={"Content-Type": mime},
        )
        return str(_json(f"PUT {path}", _expect(f"PUT {path}", made, 201))["id"])

    def create_document(self, title: str, folder: str | None) -> str:
        return self._put(self._folder(folder), f"{title}.docx", _docx(""), DOCX)

    def write(self, document: str, text: str) -> None:
        path = f"/drives/{self._drive_id()}/items/{document}/content"
        # Replacing an existing item's content answers 200 (driveitem-put-content).
        put = self._graph("PUT", path, content=_docx(text), headers={"Content-Type": DOCX})
        _expect(f"PUT {path}", put, 200)

    def download(self, document: str) -> bytes:
        """`content` answers 302 to a pre-authenticated URL, fetched without the bearer (driveitem-get-content)."""
        path = f"/drives/{self._drive_id()}/items/{document}/content"
        redirected = _expect(f"GET {path}", self._graph("GET", path), 302)
        return ok("download", self._http.get(redirected.headers["location"])).content

    def read_text(self, document: str) -> str:
        item = self._item(document)
        content = self.download(document)
        mime = str(item["file"]["mimeType"]) if "file" in item and "mimeType" in item["file"] else ""
        return _docx_text(content) if mime == DOCX else content.decode()

    def rename(self, document: str, title: str) -> None:
        name = str(self._item(document)["name"])
        ending = name[len(_title(name)) :]
        path = f"/drives/{self._drive_id()}/items/{document}"
        _json(f"PATCH {path}", self._graph("PATCH", path, json={"name": f"{title}{ending}"}))

    def move(self, document: str, folder: str) -> None:
        path = f"/drives/{self._drive_id()}/items/{document}"
        _json(f"PATCH {path}", self._graph("PATCH", path, json={"parentReference": {"id": self._folder(folder)}}))

    def share(self, document: str, email: str, role: AccessRole) -> None:
        if role not in ROLES:
            raise NotImplementedError(f"driveItem invite grants read, write or owner; Graph v1.0 has no {role.value}")
        path = f"/drives/{self._drive_id()}/items/{document}/invite"
        body = {"recipients": [{"email": email}], "roles": [ROLES[role]], "requireSignIn": True,
                "sendInvitation": False}  # fmt: skip
        _json(f"POST {path}", self._graph("POST", path, json=body))

    def trash(self, document: str) -> None:
        path = f"/drives/{self._drive_id()}/items/{document}"
        # Delete moves the item to the recycle bin and answers 204 (driveitem-delete).
        _expect(f"DELETE {path}", self._graph("DELETE", path), 204)

    def comments(self, document: str) -> list[CommentSeen]:
        raise NotImplementedError(MicrosoftDriver.absent["documents.comments"])

    def document_pages(self, page_size: int) -> list[list[str]]:
        return [
            [str(i["id"]) for i in page if _given(i, "file") and not _deleted(i)] for page in self._items(page_size)
        ]

    def upload(self, name: str, content: bytes, mime_type: str) -> str:
        return self._put(self._folder(None), name, content, mime_type)


# ---------------------------------------------------------------------------------------------- faults

FaultCall = Callable[[httpx.Client, Directory], httpx.Response]


def _app_token(http: httpx.Client, tenant: Directory, scope: str) -> dict[str, str]:
    return {"Authorization": f"Bearer {SignIn(http, tenant).app(scope)}"}


def _users(http: httpx.Client, tenant: Directory) -> httpx.Response:
    return http.get(f"{GRAPH}/users", params={"$top": 10}, headers=_app_token(http, tenant, GRAPH_DEFAULT))


def _team(http: httpx.Client, tenant: Directory) -> httpx.Response:
    return http.get(f"{GRAPH}/teams/{tenant.team_id}", headers=_app_token(http, tenant, GRAPH_DEFAULT))


def _send(http: httpx.Client, tenant: Directory) -> httpx.Response:
    return http.post(
        f"{CONNECTOR}v3/conversations/{quote(tenant.general_channel_id, safe='')}/activities",
        json={"type": "message", "text": "A conformance send"},
        headers=_app_token(http, tenant, BOT_DEFAULT),
    )


def _raised(call: FaultCall) -> Callable[[Api, WorldView], BaseException | None]:
    """The call, raising as an httpx client raises a refusal (`raise_for_status`); what it raised."""

    def trigger(api: Api, world: WorldView) -> BaseException | None:
        with api.http() as http:
            try:
                call(http, directory(world.name)).raise_for_status()
            except httpx.HTTPStatusError as refused:
                return refused
        return None

    return trigger


def _succeeds(call: FaultCall) -> Callable[[Api, WorldView], None]:
    def then(api: Api, world: WorldView) -> None:
        with api.http() as http:
            call(http, directory(world.name)).raise_for_status()

    return then


def _sent_id(api: Api, world: WorldView) -> BaseException | None:
    """A bot keeping the id of what it sent, as the connector documents a send answering it (`ResourceResponse`)."""
    with api.http() as http:
        answered = _send(http, directory(world.name))
        try:
            answered.raise_for_status()
            str(answered.json()["id"])
        except (httpx.HTTPStatusError, KeyError) as failed:
            return failed
    return None


def _id_kept(api: Api, world: WorldView) -> None:
    failed = _sent_id(api, world)
    if failed is not None:
        raise failed


def _fault(call: str, answer: Mapping[str, object], **more: object) -> Mapping[str, object]:
    return {"faults": [{"call": call, "answer": dict(answer), **more}]}


# ---------------------------------------------------------------------------------------------- the driver


class MicrosoftDriver(Driver):
    provider: ClassVar[str] = PROVIDER
    session: ClassVar[type[Session]] = MicrosoftSession
    absent: ClassVar[Mapping[str, str]] = {
        "accounts.bot": "Graph lists users at /users and an app or bot is a service principal, never a user, so no "
        "account the listing answers is a bot (https://learn.microsoft.com/en-us/graph/api/resources/user)",
        "messaging.topic": "a Teams channel has a displayName and a description and no topic; only a group chat "
        "has a topic (https://learn.microsoft.com/en-us/graph/api/resources/channel)",
        "documents.comments": "Microsoft Graph v1.0 has no API for a file's comments: a driveItem has no comments "
        "relationship (https://learn.microsoft.com/en-us/graph/api/resources/driveitem), as the provider's README says",
    }
    id_formats: ClassVar[Mapping[IdKind, re.Pattern[str]]] = {
        # A user's id is its Entra object id, a GUID (https://learn.microsoft.com/en-us/graph/api/resources/user,
        # `id`: "The unique identifier for the user"; every example is a lowercase GUID).
        IdKind.PERSON: re.compile(r"^[0-9a-f]{8}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{12}$"),
        # A channel is `19:…@thread.tacv2` (or the older `@thread.skype`); a 1:1 the bot creates is the connector's
        # `a:…` id, and Graph's `19:{user}_{other}@unq.gbl.spaces`; a group chat `19:…@thread.v2` (CLAIMS.md "an
        # `a:` id"; https://learn.microsoft.com/en-us/graph/api/resources/channel and .../resources/chat).
        IdKind.CHANNEL: re.compile(
            r"^(19:[0-9A-Za-z_-]+@thread\.(tacv2|skype|v2)|19:[0-9a-f-]+_[0-9a-f-]+@unq\.gbl\.spaces|a:[0-9A-Za-z_-]+)$"
        ),
        # A Teams message id is the moment it was sent in epoch milliseconds, 13 digits today
        # (https://learn.microsoft.com/en-us/graph/api/resources/chatmessage examples; the parentMessageId of a
        # deep link "is the timestamp" of the root, https://learn.microsoft.com/en-us/microsoftteams/platform/concepts/build-and-test/deep-link-teams).
        IdKind.MESSAGE: re.compile(r"^\d{13}$"),
        # A OneDrive for Business / SharePoint driveItem id is 34 characters, `01` and 32 uppercase base-32
        # characters (https://learn.microsoft.com/en-us/graph/api/resources/driveitem examples; the provider's
        # `item_id` keeps "Graph's 34-character shape").
        IdKind.DOCUMENT: re.compile(r"^01[A-Z0-9]{32}$"),
    }
    page_floor: ClassVar[Mapping[str, int]] = {
        # `$top` is a positive page size on every Graph list (https://learn.microsoft.com/en-us/graph/query-parameters#top);
        # users take 1 to 999, channel messages 1 to 50 (list-channel-messages), driveItem collections 1 and up.
        "people": 1,
        "history": 1,
        "documents": 1,
    }
    unknown_refusal: ClassVar[tuple[int, str]] = (401, "InvalidAuthenticationToken")
    """Graph refuses an access token it cannot validate 401 with code `InvalidAuthenticationToken`
    (https://learn.microsoft.com/en-us/graph/errors; https://learn.microsoft.com/en-us/graph/resolve-auth-errors)."""

    def __init__(self) -> None:
        self._people: dict[str, dict[str, Mapping[str, Any]]] = {}
        """Each world's people by key, by its name (the tag): whom a person's session signs in as."""

    def world(self, seed: dict[str, object], tag: str, *, logins: Mapping[str, str] | None = None) -> CreateWorld:
        if logins:
            raise NotImplementedError(
                "Microsoft accounts have a login of their own (the userPrincipalName), but the Microsoft seed has no "
                "field naming one: the provider derives it from the email's local part (seed.py `_graph_user`)"
            )
        merged = dict(seed) | {"name": tag}
        self._people[tag] = {str(p["key"]): p for p in _objects(seed, "people")}
        tenant = directory(tag)
        keys = [tenant.tenant_id, tenant.tenant_domain, tenant.sharepoint_host.split(".")[0]]
        return CreateWorld(
            seed=Seed.model_validate(merged),
            claims=Claims(keys=keys, tokens=[tenant.bot_app_id, tenant.bot_app_secret]),
        )

    def connect(self, api: Api, world: WorldView, *, person: str | None = None) -> Session:
        tenant = directory(world.name)
        if person is None:
            return MicrosoftSession(api, tenant, None)
        login = self.credential_of(world, person)
        if login is None:
            raise NotImplementedError(f"{person} is another app's bot, not a user of the tenant, and cannot sign in")
        return MicrosoftSession(api, tenant, login)

    def credential_of(self, world: WorldView, person: str | None) -> str | None:
        """The bot app's client secret for the agent; for a person, the sign-in name their authorization code flow
        names (`login_hint`), from which the platform mints their tokens. A bot account is no user and has none."""
        if person is None:
            return directory(world.name).bot_app_secret
        found = self._people[world.name][person]
        if "account" in found and found["account"] == "bot":
            return None
        return str(found["email"])

    def faults(self) -> list[FaultCase]:
        # No Microsoft client library is a dev dependency (no msgraph-sdk, no msal, no botbuilder): each fault is
        # triggered with httpx as the REST references write the call, and `typed` is None.
        return [
            FaultCase(
                name="rate_limited",
                fragment=_fault("GET /v1.0/users", {"kind": "rate_limited", "retry_after": "PT2S"}),
                trigger=_raised(_users),
                typed=None,
                status=429,
                holds="TooManyRequests",
                then=_succeeds(_users),
            ),
            FaultCase(
                name="refused_graph",
                fragment=_fault("GET /v1.0/teams", {"kind": "refused", "error": "InvalidAuthenticationToken"}, times=2),
                trigger=_raised(_team),
                typed=None,
                status=401,
                holds="InvalidAuthenticationToken",
                uses=2,
                then=_succeeds(_team),
            ),
            FaultCase(
                name="refused_connector",
                fragment=_fault("POST /teams/v3/conversations", {"kind": "refused", "error": "generalException"}),
                trigger=_raised(_send),
                typed=None,
                status=500,
                holds="generalException",
                then=_succeeds(_send),
            ),
            FaultCase(
                name="refused_only_rich",
                fragment=_fault(
                    "POST /teams/v3/conversations", {"kind": "refused", "error": "ConversationNotFound"}, only_rich=True
                ),
                trigger=_raised(_send),
                typed=None,
                status=404,
                holds="ConversationNotFound",
                then=_succeeds(_send),
            ),
            FaultCase(
                name="without_id",
                fragment=_fault("POST /teams/v3/conversations", {"kind": "without_id"}),
                trigger=_sent_id,
                typed=None,
                status=201,
                holds="{}",
                then=_id_kept,
            ),
        ]


def _objects(seed: Mapping[str, object], key: str) -> list[Mapping[str, Any]]:
    found = seed[key] if key in seed else []
    if not isinstance(found, list):
        raise TypeError(f"the seed's {key} is not a list")
    return [x for x in found if isinstance(x, Mapping)]


DRIVER = MicrosoftDriver()
